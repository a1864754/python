from docx import Document  # 导入库

data = []
for a in range(1, 29):
    print("开始第", a, "份")

    a = str(a)

    path = "C:\\Users\\vayne\\Desktop\\素质报告\\至诚书院大学生综合素质教育考核表(" + a + ").docx"

    document = Document(path)  # 读入文件
    tables = document.tables  # 获取文件中的表格集

    for table in tables[:]:
        for a, row in enumerate(table.rows[:]):  # 读每行
            row_content = []
            for cell in row.cells[:]:  # 读一行中的所有单元格
                c = cell.text
                row_content.append(c)
            print(row_content)  # 以列表形式导出每一行数据
            data.append(row_content)
# print(data)


import xlwt

workbook = xlwt.Workbook(encoding='utf-8')
worksheet = workbook.add_sheet('My Worksheet', cell_overwrite_ok=True)
a = 1
x = 1
y = 4
m = 3
n = 1
print(data)
while a < 29:
    print(a)
    data3 = data[x][y]
    data4 = data[m][n]
    print(data3)
    print(data4)

    worksheet.write(a, 0, data3)
    worksheet.write(a, 1, data4)
    a += 1
    x += 8
    m += 8
# 写入excel,1,4 3,1
# 参数对应 行, 列, 值

# 保存
workbook.save('Excel_test.xls')
